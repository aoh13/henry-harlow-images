"""
Mona Steel quotation as an editable Excel workbook — same layout as the PDF
(build_monasteel.py), with live formulas for weights and totals.

Usage:
    python build_monasteel_xlsx.py config.json OUT.xlsx
"""
import datetime as dt
import json
import os
import sys

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.drawing.spreadsheet_drawing import OneCellAnchor, AnchorMarker
from openpyxl.drawing.xdr import XDRPositiveSize2D
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.units import pixels_to_EMU

HERE = os.path.dirname(os.path.abspath(__file__))
LOGO = os.path.join(HERE, '..', 'assets', 'monasteel_logo.png')

CHARCOAL = '2D2D2D'
NEARBLK = '1A1A1A'
MIDGREY = '666666'
LINEGREY = 'CCCCCC'
ROWGREY = 'F2F2F2'
GRPGREY = 'E8E8E8'
FONT = 'Arial'

ITEM_ROWS = 10          # fixed item rows in the table (blank ones stay empty)

US_TERMS = [
    ('Shipping Terms',  ''),
    ('Payment Terms',   'Net 90 days from date of Bill of Lading.'),
    ('Quote Validity',  'This quotation is valid for 30 days from the date of issue.'),
    ('Lead Time',       'To be confirmed upon receipt of purchase order.'),
    ('Material',        'As specified per line item.'),
    ('Tolerances',      'Dimensional tolerances per ASTM A6/A6M. Weights are theoretical; '
                        'certified scale weights govern at shipment.'),
    ('Partial Orders',  'This quotation is based on the full quantity listed. Partial orders are subject to re-quote.'),
    ('Price Adjustment','Prices are subject to adjustment if steel costs change by more than 5%.'),
    ('Bank Details',    ''),
]


def f(size=9, bold=False, color=NEARBLK):
    return Font(name=FONT, size=size, bold=bold, color=color)


def side(color, style='thin'):
    return Side(style=style, color=color)


def build(cfg, out):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Quotation'
    ws.sheet_view.showGridLines = False

    widths = {'A': 5, 'B': 12, 'C': 12, 'D': 9, 'E': 8, 'F': 8, 'G': 7, 'H': 12, 'I': 12, 'J': 14}
    for c, w in widths.items():
        ws.column_dimensions[c].width = w
    cols = 'ABCDEFGHIJ'

    def merge(rng, value=None, font=None, align=None, fill=None):
        ws.merge_cells(rng)
        cell = ws[rng.split(':')[0]]
        if value is not None:
            cell.value = value
        if font: cell.font = font
        if align: cell.alignment = align
        if fill:
            for row in ws[rng]:
                for c in row:
                    c.fill = fill
        return cell

    def rule(row, color=CHARCOAL, style='medium', first='A', last='J'):
        for c in cols[cols.index(first):cols.index(last) + 1]:
            cell = ws[f'{c}{row}']
            cell.border = Border(bottom=side(color, style))

    # ── Logo + top rule ──────────────────────────────────────────────────────
    ws.row_dimensions[1].height = 18
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 8
    img = XLImage(LOGO)
    w_px = 230
    h_px = round(w_px * img.height / img.width)
    img.width, img.height = w_px, h_px
    # right-align the logo to column J's right edge: anchor in H with an offset
    hij_px = sum(int(widths[c] * 7 + 5) for c in 'HIJ')
    img.anchor = OneCellAnchor(
        _from=AnchorMarker(col=7, colOff=pixels_to_EMU(max(hij_px - w_px - 4, 0)),
                           row=1, rowOff=pixels_to_EMU(4)),
        ext=XDRPositiveSize2D(pixels_to_EMU(w_px), pixels_to_EMU(h_px)))
    ws.add_image(img)
    rule(3)

    # ── Title ────────────────────────────────────────────────────────────────
    ws.row_dimensions[5].height = 24
    merge('A5:J5', 'QUOTATION', f(15, True), Alignment(horizontal='center', vertical='center'))

    # ── Customer (left) / quote info (right) ─────────────────────────────────
    ws['A7'] = 'Customer:'; ws['A7'].font = f(9, True)
    merge('A8:E8', cfg.get('customer', ''), f(12, True))
    ws['A10'] = 'Attn:'; ws['A10'].font = f(9, True)
    merge('A11:E11', cfg.get('attn', ''), f(9))

    qdate = dt.datetime.strptime(cfg['offer_date'], '%B %d, %Y')
    info = [
        (7,  'Page',          '1 / 1', '": "@'),
        (8,  'Quote Date',    qdate, '": "mmmm d, yyyy'),
        (9,  'Quote No.',     cfg['offer_no'], '": "@'),
        (10, 'Valid Through', '=I8+30', '": "mmmm d, yyyy'),
        (11, 'Prepared By',   cfg.get('prepared_by', ''), '": "@'),
    ]
    for r, label, val, fmt in info:
        ws[f'H{r}'] = label
        ws[f'H{r}'].font = f(9, True, MIDGREY)
        c = merge(f'I{r}:J{r}', val, f(9), Alignment(horizontal='left'))
        if fmt:
            c.number_format = fmt

    # ── Opening line ─────────────────────────────────────────────────────────
    merge('A13:J13', 'Thank you for your inquiry. We are pleased to quote the following material:', f(9))

    # ── Products table ───────────────────────────────────────────────────────
    hdr = 15
    headers = ['No.', 'Size', 'Shape', 'Grade', 'Weight\n(lb/ft)', 'Length\n(ft)',
               'Qty\n(pcs)', 'Total Wt.\n(lbs)', 'Unit Price\n(USD/pc)', 'Extended\n(USD)']
    grid = Border(left=side(LINEGREY), right=side(LINEGREY), top=side(LINEGREY), bottom=side(LINEGREY))
    ws.row_dimensions[hdr].height = 28
    for c, h in zip(cols, headers):
        cell = ws[f'{c}{hdr}']
        cell.value = h
        cell.font = f(8, True, 'FFFFFF')
        cell.fill = PatternFill('solid', fgColor='333333')
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        cell.border = grid

    items = cfg['line_items']
    first = hdr + 1
    last = hdr + ITEM_ROWS
    for i in range(ITEM_ROWS):
        r = first + i
        it = items[i] if i < len(items) else None
        ws.row_dimensions[r].height = 15
        values = {
            'A': f'=IF(B{r}="","",ROW()-{hdr})',
            'B': it['size'] if it else None,
            'C': it['shape'] if it else None,
            'D': (it.get('grade') or None) if it else None,
            'E': it['lb_ft'] if it else None,
            'F': it['length_ft'] if it else None,
            'G': it['qty'] if it else None,
            'H': f'=IF(OR(E{r}="",F{r}="",G{r}=""),"",E{r}*F{r}*G{r})',
            'I': (it.get('unit_price') or 0) if it else None,
            'J': f'=IF(B{r}="","",G{r}*N(I{r}))',
        }
        for c in cols:
            cell = ws[f'{c}{r}']
            cell.value = values[c]
            cell.font = f(8)
            cell.border = grid
            cell.alignment = Alignment(horizontal='left' if c == 'B' else ('right' if c in 'IJ' else 'center'),
                                       vertical='center')
            if i % 2 == 1:
                cell.fill = PatternFill('solid', fgColor=ROWGREY)
        ws[f'E{r}'].number_format = '0.##'
        ws[f'F{r}'].number_format = '0"\'"'
        ws[f'G{r}'].number_format = '#,##0'
        ws[f'H{r}'].number_format = '#,##0'
        ws[f'I{r}'].number_format = '#,##0.00;-#,##0.00;"—"'
        ws[f'J{r}'].number_format = '#,##0.00;-#,##0.00;"—"'

    # subtotal
    sr = last + 1
    gfill = PatternFill('solid', fgColor=GRPGREY)
    ws.merge_cells(f'A{sr}:H{sr}')
    for c in cols:
        cell = ws[f'{c}{sr}']
        cell.fill = gfill
        cell.border = Border(top=side(MIDGREY), bottom=side(MIDGREY),
                             left=side(LINEGREY), right=side(LINEGREY))
    ws[f'I{sr}'] = 'Subtotal'
    ws[f'I{sr}'].font = f(8, True)
    ws[f'I{sr}'].alignment = Alignment(horizontal='right')
    ws[f'J{sr}'] = f'=SUM(J{first}:J{last})'
    ws[f'J{sr}'].font = f(8, True)
    ws[f'J{sr}'].number_format = '#,##0.00;-#,##0.00;"—"'

    # total
    tr = sr + 1
    ws.row_dimensions[tr].height = 18
    bfill = PatternFill('solid', fgColor=NEARBLK)
    merge(f'A{tr}:I{tr}', 'TOTAL', f(9, True, 'FFFFFF'), Alignment(horizontal='center', vertical='center'), bfill)
    ws[f'J{tr}'] = f'=J{sr}'
    ws[f'J{tr}'].font = f(9, True, 'FFFFFF')
    ws[f'J{tr}'].fill = bfill
    ws[f'J{tr}'].alignment = Alignment(horizontal='right', vertical='center')
    ws[f'J{tr}'].number_format = '"USD "#,##0.00;"USD "-#,##0.00;"USD —"'
    for c in cols:
        ws[f'{c}{tr}'].border = Border(top=side(CHARCOAL, 'medium'))

    # ── Terms ────────────────────────────────────────────────────────────────
    r = tr + 2
    ws[f'A{r}'] = 'TERMS AND CONDITIONS'
    ws[f'A{r}'].font = f(10, True)
    rule(r, LINEGREY, 'thin')
    overrides = cfg.get('terms', {})
    for k, v in US_TERMS:
        r += 1
        v = overrides.get(k, v)
        merge(f'A{r}:B{r}', k, f(9, True), Alignment(vertical='top'))
        merge(f'C{r}:J{r}', v, f(9), Alignment(vertical='top', wrap_text=True))
        ws.row_dimensions[r].height = 15
        rule(r, 'EEEEEE', 'thin')

    # ── Closing + signatures ─────────────────────────────────────────────────
    r += 2
    merge(f'A{r}:J{r}', 'We appreciate the opportunity to quote and look forward to receiving your order.', f(9))
    r += 2
    ws[f'A{r}'] = 'Authorized By:'; ws[f'A{r}'].font = f(9, True)
    ws[f'F{r}'] = 'Customer Acceptance (Signature & Date):'; ws[f'F{r}'].font = f(9, True)
    r += 1
    ws.row_dimensions[r].height = 34
    rule(r, MIDGREY, 'thin', 'A', 'D')
    rule(r, MIDGREY, 'thin', 'F', 'J')
    r += 1
    ws[f'A{r}'] = cfg['company']['name']; ws[f'A{r}'].font = f(8)

    # ── Footer (in-sheet, charcoal rule above) ───────────────────────────────
    r += 3
    for c in cols:
        ws[f'{c}{r}'].border = Border(top=side(CHARCOAL, 'medium'))
    l1, l2 = cfg['company']['footer']
    ws[f'A{r}'] = l1; ws[f'A{r}'].font = f(7, color=MIDGREY)
    ws[f'J{r}'] = 'Page 1'; ws[f'J{r}'].font = f(7, color=MIDGREY)
    ws[f'J{r}'].alignment = Alignment(horizontal='right')
    ws[f'A{r+1}'] = l2; ws[f'A{r+1}'].font = f(7, color=MIDGREY)

    # ── Print setup: US Letter, one page ─────────────────────────────────────
    ws.print_area = f'A1:J{r+1}'
    ws.page_setup.paperSize = ws.PAPERSIZE_LETTER
    ws.page_setup.orientation = 'portrait'
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 1
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.6
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True

    # ── How-to sheet (not printed with the quote) ────────────────────────────
    h = wb.create_sheet('How to use')
    h.column_dimensions['A'].width = 24
    h.column_dimensions['B'].width = 90
    rows = [
        ('Editable cells', ''),
        ('Customer / Attn', 'Quotation!A8 and A11'),
        ('Quote Date', 'Quotation!I8 — Valid Through (I10) is Quote Date + 30 days automatically'),
        ('Quote No. / Prepared By', 'Quotation!I9 and I11'),
        ('Line items', f'Quotation!B{first}:G{last} (Size, Shape, Grade, lb/ft, Length ft, Qty) and '
                       f'I{first}:I{last} (Unit Price, USD per piece). Leave a row\'s Size blank to keep it empty.'),
        ('Calculated — do not type', f'No. (A), Total Wt. lbs (H = lb/ft × ft × qty), Extended (J = qty × unit price), '
                                     f'Subtotal (J{sr}) and TOTAL (J{tr})'),
        ('Terms', f'Quotation!C{tr+3}:C{tr+11} — Shipping Terms and Bank Details are left blank to fill in'),
        ('', ''),
        ('Weights', 'lb/ft for W-shapes is the number after the "x" (W16x67 → 67 lb/ft); '
                    'example rows use values given by Onur, all lengths 40 ft.'),
    ]
    for i, (a, b) in enumerate(rows, 1):
        h[f'A{i}'] = a; h[f'B{i}'] = b
        h[f'A{i}'].font = f(10, True if i == 1 or a else False)
        h[f'B{i}'].font = f(10)
        h[f'B{i}'].alignment = Alignment(wrap_text=True, vertical='top')
        h[f'A{i}'].alignment = Alignment(vertical='top')

    wb.active = 0
    wb.save(out)
    print(f'OK | {out}')


if __name__ == '__main__':
    with open(sys.argv[1], encoding='utf-8') as fh:
        build(json.load(fh), sys.argv[2])
